import io
import os
import zipfile
import unittest
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
import docx
import pptx
from pptx.util import Inches, Pt
from app import app

class ConversionTestCase(unittest.TestCase):
    def setUp(self):
        self.app = app.test_client()
        self.app.testing = True

    def create_dummy_pdf(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(72, 700, "SwiftDocs Test Document")
        c.setFont("Helvetica", 12)
        c.drawString(72, 670, "This is a test paragraph for converting PDF.")
        c.drawString(72, 650, "Second line of test content: Col1\tCol2\tCol3")
        c.showPage()
        c.save()
        buf.seek(0)
        return buf

    def create_dummy_docx(self):
        buf = io.BytesIO()
        doc = docx.Document()
        doc.add_heading("SwiftDocs Word Sample", level=1)
        doc.add_paragraph("This is paragraph text in Word.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Data 1"
        table.cell(1, 1).text = "Data 2"
        doc.save(buf)
        buf.seek(0)
        return buf

    def create_dummy_pptx(self):
        buf = io.BytesIO()
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "SwiftDocs Presentation Title"
        subtitle.text = "Subtitle slide content"
        prs.save(buf)
        buf.seek(0)
        return buf

    def test_pdf_to_word_produces_valid_docx(self):
        pdf_buf = self.create_dummy_pdf()
        response = self.app.post(
            '/api/convert/pdf-to-word',
            data={'file': (pdf_buf, 'sample.pdf')},
            content_type='multipart/form-data'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

        docx_bytes = io.BytesIO(response.data)
        with zipfile.ZipFile(docx_bytes, 'r') as z:
            namelist = z.namelist()
            self.assertIn('word/document.xml', namelist)
            self.assertIn('[Content_Types].xml', namelist)
            doc_xml = z.read('word/document.xml').decode('utf-8')
            self.assertTrue(len(doc_xml) > 50)

    def test_pdf_to_powerpoint_produces_valid_pptx(self):
        pdf_buf = self.create_dummy_pdf()
        response = self.app.post(
            '/api/convert/pdf-to-powerpoint',
            data={'file': (pdf_buf, 'sample.pdf')},
            content_type='multipart/form-data'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        )

        pptx_bytes = io.BytesIO(response.data)
        with zipfile.ZipFile(pptx_bytes, 'r') as z:
            namelist = z.namelist()
            self.assertIn('ppt/presentation.xml', namelist)
            self.assertIn('ppt/slides/slide1.xml', namelist)
            self.assertIn('[Content_Types].xml', namelist)

    def test_pdf_to_excel_produces_valid_xlsx(self):
        pdf_buf = self.create_dummy_pdf()
        response = self.app.post(
            '/api/convert/pdf-to-excel',
            data={'file': (pdf_buf, 'sample.pdf')},
            content_type='multipart/form-data'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        xlsx_bytes = io.BytesIO(response.data)
        with zipfile.ZipFile(xlsx_bytes, 'r') as z:
            namelist = z.namelist()
            self.assertIn('xl/workbook.xml', namelist)
            self.assertIn('[Content_Types].xml', namelist)

    def test_word_to_pdf_produces_valid_pdf(self):
        docx_buf = self.create_dummy_docx()
        response = self.app.post(
            '/api/convert/word-to-pdf',
            data={'file': (docx_buf, 'sample.docx')},
            content_type='multipart/form-data'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/pdf')
        self.assertTrue(response.data.startswith(b'%PDF-'))

    def test_powerpoint_to_pdf_produces_valid_pdf(self):
        pptx_buf = self.create_dummy_pptx()
        response = self.app.post(
            '/api/convert/powerpoint-to-pdf',
            data={'file': (pptx_buf, 'sample.pptx')},
            content_type='multipart/form-data'
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/pdf')
        self.assertTrue(response.data.startswith(b'%PDF-'))

if __name__ == '__main__':
    unittest.main()
